"""StudentHub - helpers for the corrected practical reports (python-docx).

The layout follows the existing ITUE203 reports of this repository (A4, 2 cm margins, Calibri 11, navy headings, the same
cover block and student table, numbered sections, captions under figures, page footer). Code excerpts are read from the
real source files and figures are genuine screenshots from the evidence/ folders, so nothing in a report is typed by hand
that could differ from the project.
"""
import pathlib, re
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

ROOT = pathlib.Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "submission-reports"
NAVY = RGBColor(0x1E, 0x3A, 0x8A)
STUDENT_TABLE = [["Student Name", "Shreyas"], ["Roll Number", "25CS023"], ["Student ID", "25CS023"],
                 ["Email", "25cs023@charusat.edu.in"], ["Course", "ITUE203 – Web Development Frameworks"],
                 ["Programme", "B.Tech Computer Science & Engineering"], ["Semester / Academic Year", "3 / 2026-27"],
                 ["Date of Submission", ""], ["Faculty Signature", ""]]


def report_file(number, supersedes):
    """file name of a report: '..._corrected.docx' when it replaces an earlier report, otherwise the normal name"""
    return f"PRAC_{number}_ITUE203_25CS023" + ("_corrected" if supersedes else "") + ".docx"


def test_run():
    """the suites of the saved complete run tests/run_all_output.txt: ([(suite, passed, failed, seconds)], total_passed, total_failed)"""
    suites, total = [], None
    for line in (ROOT / "tests" / "run_all_output.txt").read_text(encoding="utf-8").splitlines():
        m = re.match(r"(.+?)\s+(\d+)\s+(\d+)\s+(\d+)s\s*$", line)
        t = re.match(r"TOTAL\s+(\d+)\s+(\d+)\s*$", line)
        if m:
            suites.append((m[1].strip(), int(m[2]), int(m[3]), int(m[4])))
        elif t:
            total = int(t[1]), int(t[2])
    assert suites and total and total == (sum(x[1] for x in suites), sum(x[2] for x in suites)), "incomplete run_all_output.txt"
    return suites, total[0], total[1]


class Report:
    def __init__(self, number, title, tech_line, supersedes):
        self.number, self.title, self.tech_line, self.supersedes = number, title, tech_line, supersedes
        self.figure_no = 0
        doc = self.doc = Document()
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Emu(7560310), Emu(10692130)            # A4
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Cm(2))
        normal = doc.styles["Normal"]
        normal.font.name, normal.font.size = "Calibri", Pt(11)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.12
        for name, size in (("Heading 1", 15), ("Heading 2", 12.5), ("Title", 22)):
            st = doc.styles[name]
            st.font.name, st.font.size, st.font.color.rgb = "Calibri", Pt(size), NAVY
        doc.styles["Heading 1"].paragraph_format.space_before = Pt(16)
        doc.styles["Heading 1"].paragraph_format.keep_with_next = True
        code = doc.styles.add_style("Code", 1)
        code.font.name, code.font.size = "Consolas", Pt(8)
        code.paragraph_format.left_indent, code.paragraph_format.first_line_indent = Cm(0.9), Cm(-0.9)
        code.paragraph_format.space_after, code.paragraph_format.line_spacing = Pt(0), 1.0
        cap = doc.styles["Caption"]
        cap.font.size, cap.font.italic, cap.font.color.rgb = Pt(9.5), True, RGBColor(0x37, 0x41, 0x51)
        self._cover()

    # ------------------------------------------------------------ text helpers
    def _runs(self, p, text, bold=False, size=None, color=None):
        """`code` spans become Consolas, **bold** spans bold, *italic* spans italic"""
        for part in re.split(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*\s][^*]*\*)", text):
            if not part:
                continue
            if part.startswith("`"):
                r = p.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9.5)
            elif part.startswith("**"):
                r = p.add_run(part[2:-2]); r.bold = True
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                r = p.add_run(part[1:-1]); r.italic = True; r.bold = bold
            else:
                r = p.add_run(part); r.bold = bold
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color

    def para(self, text="", bold=False, size=None, align=None, color=None, after=None, style=None):
        p = self.doc.add_paragraph(style=style)
        self._runs(p, text, bold=bold, size=size, color=color)
        if align is not None:
            p.alignment = align
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        return p

    def heading(self, text, level=1):
        return self.doc.add_heading(text, level)

    def bullets(self, items):
        for item in items:
            self._runs(self.doc.add_paragraph(style="List Bullet"), item)

    def numbered(self, items):
        for item in items:
            self._runs(self.doc.add_paragraph(style="List Number"), item)

    @staticmethod
    def _shade(cell, color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"), shd.set(qn("w:color"), "auto"), shd.set(qn("w:fill"), color)
        tcPr.append(shd)

    def table(self, rows, widths, header=True, size=9.5):
        t = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, row in enumerate(rows):
            for j, value in enumerate(row):
                cell = t.cell(i, j)
                cell.width = Cm(widths[j])
                cell.paragraphs[0].paragraph_format.space_after = Pt(1)
                self._runs(cell.paragraphs[0], str(value), bold=header and i == 0, size=size)
                if header and i == 0:
                    self._shade(cell, "DCE4F2")
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def note(self, text, fill="FFF7E6"):
        """a shaded one-cell box (used for the 'corrected report' notice and important remarks)"""
        t = self.doc.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        cell = t.cell(0, 0)
        self._shade(cell, fill)
        self._runs(cell.paragraphs[0], text, size=10)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def excerpt(self, file, start, end, title, max_lines=45):
        """the lines of a real source file from the first line containing `start` to the next containing `end`"""
        lines = (ROOT / file).read_text(encoding="utf-8").splitlines()
        a = next(i for i, l in enumerate(lines) if start in l)
        b = next(i for i in range(a, len(lines)) if end in lines[i])
        b = min(b, a + max_lines - 1)
        self.para(f"{file} (lines {a + 1}–{b + 1}) – {title}", bold=True, size=9.5, after=2)
        t = self.doc.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        cell = t.cell(0, 0)
        self._shade(cell, "F3F5F9")
        cell.paragraphs[0].style = self.doc.styles["Code"]
        for k, line in enumerate(lines[a:b + 1]):
            p = cell.paragraphs[0] if k == 0 else cell.add_paragraph(style=self.doc.styles["Code"])
            p.add_run(f"{a + 1 + k:>4}  {line}")
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def figure(self, path, caption, width=16.4):
        """a genuine screenshot (path relative to the repository root) with a numbered caption"""
        from PIL import Image
        full = ROOT / path
        self.figure_no += 1
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        px_w, px_h = Image.open(full).size
        width = min(width, 17.5 * px_w / px_h)                 # tall full-page captures: at most 17.5 cm high
        p.add_run().add_picture(str(full), width=Cm(width))
        c = self.doc.add_paragraph(f"Figure {self.figure_no}: {caption}  ({pathlib.PurePath(path).name})", style="Caption")
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------------------------------------ cover and footer
    def _cover(self):
        for text, size, bold in (("CHAROTAR UNIVERSITY OF SCIENCE AND TECHNOLOGY (CHARUSAT)", 14, True),
                                 ("Chandubhai S. Patel Institute of Technology (CSPIT)", 12, False),
                                 ("Department of Computer Science & Engineering", 12, False),
                                 ("ITUE203 – Web Development Frameworks", 12, False)):
            self.para(text, bold=bold, size=size, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
        self.para("", after=30)
        self.para(f"PRACTICAL {self.number}", bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, color=NAVY, after=8)
        t = self.doc.add_paragraph(self.title, style="Title")
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.para(self.tech_line, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
        self.para(("Corrected submission" if self.supersedes else "Report") + " – implemented in the integrated StudentHub portal (studenthub/)", bold=True, size=11,
                  align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0x0F, 0x76, 0x6E), after=26)
        self.table(STUDENT_TABLE, [6, 9], header=False, size=11)
        self.page_break()

    def save(self):
        footer = self.doc.sections[0].footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = footer.add_run(f"ITUE203 · Practical {self.number}{' (corrected)' if self.supersedes else ''} · Shreyas (25CS023) · Page ")
        r.font.size = Pt(9)
        for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
            run = footer.add_run()
            run.font.size = Pt(9)
            if kind:
                el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
            else:
                el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
            run._r.append(el)
        OUT_DIR.mkdir(exist_ok=True)
        out = OUT_DIR / report_file(self.number, self.supersedes)
        self.doc.save(str(out))
        return out


def corrected_notice(r, original, what_changed):
    """the standard first box of every corrected report"""
    r.note(f"**Corrected report.** This report replaces the original Practical {r.number} report (`{original}`, kept unchanged in the "
           f"repository). {what_changed} All screenshots in this report were taken from the canonical integrated application "
           f"`studenthub/` with Microsoft Edge on disposable test databases with fictional records; all test results are copied from "
           f"the saved test outputs of the repository.")
